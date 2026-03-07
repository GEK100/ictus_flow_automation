"""Tests for CX-06: Supported File Types and Handling.

Tests that .docx and .xlsx files are classified via text extraction
(not OCR), and unsupported types return UNSUPPORTED.

Run: python -m pytest tests/file-classifier/test_file_classifier.py -v
"""

import sys
import os
import json
import tempfile
import pytest
from pathlib import Path
from unittest.mock import patch, MagicMock

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from scripts.skills.enhanced_classifier import (
    classify_file,
    _extract_text_content,
    SUPPORTED_TYPES,
)


# ── Helpers ────────────────────────────────────────────────────────

MOCK_ROUTING_RULES = {
    'confidence_thresholds': {
        'auto_route': 0.90,
        'spot_check': 0.70,
        'hold': 0.00,
    },
    'classification_map': {
        'INVOICE': {'handler': 'invoice_handler'},
        'OTHER': {'handler': None},
    },
}


def _patch_classify(monkeypatch):
    """Patch external dependencies for classify_file."""
    monkeypatch.setattr(
        'scripts.skills.enhanced_classifier.load_routing_rules',
        lambda: MOCK_ROUTING_RULES,
    )
    monkeypatch.setattr(
        'scripts.skills.enhanced_classifier.load_client_config',
        lambda code: {'tracking_sheet_id': None, 'folders': {'learning': 'x'}},
    )
    monkeypatch.setattr(
        'scripts.skills.enhanced_classifier.load_classification_corrections',
        lambda code: [],
    )
    monkeypatch.setattr(
        'scripts.skills.enhanced_classifier.load_base_prompt',
        lambda name: 'Classify this document.',
    )


# ── Test 1: .docx classified without OCR ──────────────────────────

class TestDocxClassifiedWithoutOcr:

    def test_docx_classified_without_ocr(self, monkeypatch, tmp_path):
        """A .docx file extracts text via python-docx and classifies
        without triggering OCR."""
        _patch_classify(monkeypatch)

        # Create a real .docx file
        from docx import Document
        doc = Document()
        doc.add_paragraph('INVOICE')
        doc.add_paragraph('Invoice Number: INV-001')
        doc.add_paragraph('Total: £1,500.00')
        docx_path = str(tmp_path / 'test_invoice.docx')
        doc.save(docx_path)

        # Mock the Claude API call to return a classification
        mock_response = {
            'classification': 'INVOICE',
            'confidence': 0.95,
            'summary': 'An invoice document',
        }
        monkeypatch.setattr(
            'scripts.skills.enhanced_classifier.claude_client.send_message_json',
            lambda **kwargs: mock_response,
        )

        result = classify_file(
            client_code='TEST',
            local_path=docx_path,
        )

        assert result['classification'] == 'INVOICE'
        assert result['confidence'] == 0.95
        assert result['ocr_applied'] is False
        assert result['status'] == 'NEW'

    def test_docx_text_extraction(self, tmp_path):
        """Verify _extract_text_content extracts text from .docx."""
        from docx import Document
        doc = Document()
        doc.add_paragraph('Purchase Order #12345')
        doc.add_paragraph('Please supply 100 units.')
        docx_path = str(tmp_path / 'test.docx')
        doc.save(docx_path)

        text = _extract_text_content(docx_path)
        assert 'Purchase Order #12345' in text
        assert '100 units' in text


# ── Test 2: .xlsx classified ──────────────────────────────────────

class TestXlsxClassified:

    def test_xlsx_classified(self, monkeypatch, tmp_path):
        """A .xlsx file extracts headers/data and classifies correctly."""
        _patch_classify(monkeypatch)

        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['Invoice No', 'Date', 'Amount', 'VAT', 'Total'])
        ws.append(['INV-001', '2026-01-15', 1000, 200, 1200])
        ws.append(['INV-002', '2026-01-16', 2000, 400, 2400])
        xlsx_path = str(tmp_path / 'invoices.xlsx')
        wb.save(xlsx_path)

        mock_response = {
            'classification': 'INVOICE',
            'confidence': 0.92,
            'summary': 'Spreadsheet of invoices',
        }
        monkeypatch.setattr(
            'scripts.skills.enhanced_classifier.claude_client.send_message_json',
            lambda **kwargs: mock_response,
        )

        result = classify_file(
            client_code='TEST',
            local_path=xlsx_path,
        )

        assert result['classification'] == 'INVOICE'
        assert result['confidence'] == 0.92
        assert result['ocr_applied'] is False

    def test_xlsx_text_extraction(self, tmp_path):
        """Verify _extract_text_content extracts data from .xlsx."""
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['Name', 'Amount'])
        ws.append(['Widget', 500])
        xlsx_path = str(tmp_path / 'test.xlsx')
        wb.save(xlsx_path)

        text = _extract_text_content(xlsx_path)
        assert 'Name' in text
        assert 'Widget' in text


# ── Test 3: Unsupported type flagged ──────────────────────────────

class TestUnsupportedTypeFlagged:

    def test_unsupported_type_flagged(self, monkeypatch, tmp_path):
        """.zip file returns UNSUPPORTED with NEEDS_CLASSIFICATION status."""
        _patch_classify(monkeypatch)

        zip_path = str(tmp_path / 'archive.zip')
        with open(zip_path, 'wb') as f:
            f.write(b'PK\x03\x04fake zip content')

        result = classify_file(
            client_code='TEST',
            local_path=zip_path,
        )

        assert result['classification'] == 'UNSUPPORTED'
        assert result['confidence'] == 0.0
        assert result['status'] == 'NEEDS_CLASSIFICATION'
        assert '.zip' in result['summary']
        assert result['ocr_applied'] is False

    def test_unsupported_exe(self, monkeypatch, tmp_path):
        """.exe file also returns UNSUPPORTED."""
        _patch_classify(monkeypatch)

        exe_path = str(tmp_path / 'program.exe')
        with open(exe_path, 'wb') as f:
            f.write(b'\x00\x00')

        result = classify_file(
            client_code='TEST',
            local_path=exe_path,
        )

        assert result['classification'] == 'UNSUPPORTED'
        assert result['status'] == 'NEEDS_CLASSIFICATION'

    def test_supported_types_dict_complete(self):
        """SUPPORTED_TYPES covers all expected extensions."""
        expected = {'.pdf', '.jpg', '.jpeg', '.png', '.docx', '.doc',
                    '.xlsx', '.xls', '.csv', '.msg', '.eml'}
        assert set(SUPPORTED_TYPES.keys()) == expected
