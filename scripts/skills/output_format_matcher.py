"""Skill 9: Output Format Matcher.

Analyses client's existing templates from 06-TEMPLATES and creates
replication specs so outputs match their structure. Column order,
formatting, fonts, margins — all captured so handlers can reproduce
the client's familiar layout.

Usage:
    python scripts/skills/output_format_matcher.py <client_code>
"""

import os
import sys
import json
import argparse
import tempfile
import logging
from datetime import datetime
from pathlib import Path

logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')
log = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from scripts.utils import claude_client, drive_client
from scripts.utils.prompt_builder import load_base_prompt, load_client_config


FORMAT_ANALYSIS_PROMPT = None


def get_analysis_prompt():
    global FORMAT_ANALYSIS_PROMPT
    if FORMAT_ANALYSIS_PROMPT is None:
        FORMAT_ANALYSIS_PROMPT = load_base_prompt('output-format-matcher.txt')
    return FORMAT_ANALYSIS_PROMPT


def extract_spreadsheet_structure(filepath):
    """Extract structure from Excel files using openpyxl.

    Returns dict with columns, formatting, formulas, sheet names.
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(filepath, data_only=False)
    sheets = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row is None or ws.max_row < 1:
            continue

        columns = []
        # Read header row (row 1)
        for col_idx in range(1, (ws.max_column or 0) + 1):
            header_cell = ws.cell(row=1, column=col_idx)
            header_value = header_cell.value
            if header_value is None:
                continue

            # Detect data type from first few data rows
            sample_values = []
            for row_idx in range(2, min(ws.max_row + 1, 12)):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.value is not None:
                    sample_values.append(cell.value)

            col_type = _infer_column_type(sample_values)

            # Check for number format
            number_format = None
            for row_idx in range(2, min(ws.max_row + 1, 6)):
                cell = ws.cell(row=row_idx, column=col_idx)
                if cell.number_format and cell.number_format != 'General':
                    number_format = cell.number_format
                    break

            # Check for formulas
            formula_pattern = None
            for row_idx in range(2, min(ws.max_row + 1, 6)):
                cell = ws.cell(row=row_idx, column=col_idx)
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    formula_pattern = cell.value
                    break

            col_info = {
                'name': str(header_value),
                'position': col_idx,
                'type': col_type,
            }
            if number_format:
                col_info['format'] = number_format
            if formula_pattern:
                col_info['formula_pattern'] = formula_pattern

            columns.append(col_info)

        # Extract conditional formatting rules
        conditional_formatting = []
        for cf_rule in ws.conditional_formatting:
            for rule in cf_rule.rules:
                cf_entry = {
                    'range': str(cf_rule),
                    'type': rule.type,
                }
                if rule.formula:
                    cf_entry['formula'] = [str(f) for f in rule.formula]
                if hasattr(rule, 'dxf') and rule.dxf:
                    if rule.dxf.fill and rule.dxf.fill.fgColor:
                        cf_entry['fill_color'] = str(rule.dxf.fill.fgColor.rgb)
                    if rule.dxf.font and rule.dxf.font.color:
                        cf_entry['font_color'] = str(rule.dxf.font.color.rgb)
                conditional_formatting.append(cf_entry)

        # Column widths
        col_widths = {}
        for col_letter, dim in ws.column_dimensions.items():
            if dim.width and dim.width != 8.43:  # skip default width
                col_widths[col_letter] = dim.width

        sheets[sheet_name] = {
            'columns': columns,
            'row_count': ws.max_row,
            'conditional_formatting': conditional_formatting,
            'column_widths': col_widths,
            'frozen_panes': str(ws.freeze_panes) if ws.freeze_panes else None,
        }

    wb.close()
    return {'type': 'spreadsheet', 'sheets': sheets}


def extract_csv_structure(filepath):
    """Extract structure from CSV files.

    Returns dict with column names and inferred types.
    """
    import csv

    with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
        reader = csv.reader(f)
        try:
            headers = next(reader)
        except StopIteration:
            return {'type': 'spreadsheet', 'sheets': {}}

        # Read sample rows for type inference
        sample_rows = []
        for i, row in enumerate(reader):
            if i >= 10:
                break
            sample_rows.append(row)

    columns = []
    for col_idx, header in enumerate(headers):
        if not header.strip():
            continue
        sample_values = []
        for row in sample_rows:
            if col_idx < len(row) and row[col_idx]:
                sample_values.append(row[col_idx])
        col_type = _infer_column_type(sample_values)
        columns.append({
            'name': header.strip(),
            'position': col_idx + 1,
            'type': col_type,
        })

    return {
        'type': 'spreadsheet',
        'sheets': {
            'Sheet1': {
                'columns': columns,
                'row_count': len(sample_rows) + 1,
                'conditional_formatting': [],
                'column_widths': {},
                'frozen_panes': None,
            }
        },
    }


def extract_docx_structure(filepath):
    """Extract structure from Word documents using python-docx.

    Returns dict with layout, styles, sections.
    """
    from docx import Document
    from docx.shared import Inches, Pt, Cm

    doc = Document(filepath)
    sections_info = []
    for section in doc.sections:
        sec = {
            'page_width': section.page_width.inches if section.page_width else None,
            'page_height': section.page_height.inches if section.page_height else None,
            'orientation': 'landscape' if section.orientation else 'portrait',
            'margins': {
                'top': section.top_margin.inches if section.top_margin else None,
                'bottom': section.bottom_margin.inches if section.bottom_margin else None,
                'left': section.left_margin.inches if section.left_margin else None,
                'right': section.right_margin.inches if section.right_margin else None,
            },
        }
        sections_info.append(sec)

    # Extract heading and paragraph styles used
    styles_used = {}
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else 'Normal'
        if style_name not in styles_used:
            font_info = {}
            if para.style and para.style.font:
                font = para.style.font
                if font.name:
                    font_info['name'] = font.name
                if font.size:
                    font_info['size_pt'] = font.size.pt
                if font.bold:
                    font_info['bold'] = True
                if font.italic:
                    font_info['italic'] = True
                if font.color and font.color.rgb:
                    font_info['color'] = str(font.color.rgb)
            styles_used[style_name] = {
                'count': 0,
                'font': font_info,
            }
        styles_used[style_name]['count'] += 1

    # Check for tables
    tables = []
    for table in doc.tables:
        table_info = {
            'rows': len(table.rows),
            'columns': len(table.columns),
            'headers': [],
        }
        if table.rows:
            for cell in table.rows[0].cells:
                table_info['headers'].append(cell.text.strip())
        tables.append(table_info)

    # Check for headers/footers
    has_header = False
    has_footer = False
    for section in doc.sections:
        if section.header and section.header.paragraphs:
            header_text = ' '.join(p.text for p in section.header.paragraphs).strip()
            if header_text:
                has_header = True
        if section.footer and section.footer.paragraphs:
            footer_text = ' '.join(p.text for p in section.footer.paragraphs).strip()
            if footer_text:
                has_footer = True

    return {
        'type': 'document',
        'sections': sections_info,
        'styles': styles_used,
        'tables': tables,
        'has_header': has_header,
        'has_footer': has_footer,
        'paragraph_count': len(doc.paragraphs),
    }


def extract_pdf_structure(filepath):
    """Extract layout structure from PDF files.

    Returns dict with page dimensions, text layout info.
    """
    try:
        import pdfplumber

        with pdfplumber.open(filepath) as pdf:
            pages_info = []
            for i, page in enumerate(pdf.pages[:5]):
                page_info = {
                    'width': page.width,
                    'height': page.height,
                    'text_length': len(page.extract_text() or ''),
                }
                # Check for tables
                tables = page.extract_tables()
                if tables:
                    page_info['tables'] = []
                    for table in tables:
                        if table and table[0]:
                            page_info['tables'].append({
                                'rows': len(table),
                                'columns': len(table[0]),
                                'headers': [str(h) for h in table[0] if h],
                            })
                pages_info.append(page_info)

            return {
                'type': 'pdf',
                'page_count': len(pdf.pages),
                'pages': pages_info,
            }
    except ImportError:
        # Fallback to PyPDF2
        from PyPDF2 import PdfReader

        reader = PdfReader(filepath)
        pages_info = []
        for i, page in enumerate(reader.pages[:5]):
            box = page.mediabox
            pages_info.append({
                'width': float(box.width),
                'height': float(box.height),
                'text_length': len(page.extract_text() or ''),
            })

        return {
            'type': 'pdf',
            'page_count': len(reader.pages),
            'pages': pages_info,
        }


def _infer_column_type(values):
    """Infer column data type from sample values."""
    if not values:
        return 'text'

    date_indicators = ['/', '-', 'jan', 'feb', 'mar', 'apr', 'may', 'jun',
                       'jul', 'aug', 'sep', 'oct', 'nov', 'dec']
    num_count = 0
    date_count = 0
    currency_count = 0

    for v in values:
        s = str(v).strip().lower()

        # Currency check
        if s.startswith(('$', '\u00a3', '\u20ac')) or s.endswith(('p', 'gbp', 'usd', 'eur')):
            currency_count += 1
            continue

        # Number check
        cleaned = s.replace(',', '').replace('%', '').strip()
        try:
            float(cleaned)
            num_count += 1
            continue
        except ValueError:
            pass

        # Date check
        if any(ind in s for ind in date_indicators):
            date_count += 1

    total = len(values)
    if currency_count > total * 0.5:
        return 'currency'
    if num_count > total * 0.5:
        return 'number'
    if date_count > total * 0.5:
        return 'date'
    return 'text'


def analyse_template(filepath, filename, client_code):
    """Extract structure from a template file and send to Sonnet for spec creation.

    Returns the structured format spec dict.
    """
    ext = Path(filepath).suffix.lower()

    # Extract raw structure based on file type
    if ext in {'.xlsx', '.xls'}:
        structure = extract_spreadsheet_structure(filepath)
    elif ext == '.csv':
        structure = extract_csv_structure(filepath)
    elif ext == '.docx':
        structure = extract_docx_structure(filepath)
    elif ext == '.pdf':
        structure = extract_pdf_structure(filepath)
    else:
        log.warning(f"Unsupported file type: {ext} for {filename}")
        return None

    # Send to Sonnet for structured replication spec
    system_prompt = get_analysis_prompt()
    prompt = (
        f"Analyse this template structure and create a replication specification.\n\n"
        f"Template filename: {filename}\n"
        f"File type: {ext}\n\n"
        f"Extracted structure:\n{json.dumps(structure, indent=2, default=str)}"
    )

    return claude_client.send_message_json(
        prompt=prompt,
        system_prompt=system_prompt,
        model_tier='MODEL_STANDARD',
        max_tokens=3000,
        client_code=client_code,
        workflow='output-format-matcher',
    )


def get_format_spec(client_code, output_type):
    """Helper for handlers: get the format spec for a given output type.

    Args:
        client_code: Client code (e.g. 'GILM').
        output_type: Output type key (e.g. 'invoice_tracker').

    Returns:
        Format spec dict or None if not found.
    """
    config = load_client_config(client_code)
    if not config:
        return None
    learning_folder = config.get('folders', {}).get('learning')
    if not learning_folder:
        return None

    formats_data = drive_client.download_json(learning_folder, 'output-formats.json')
    if not formats_data:
        return None

    return formats_data.get('formats', {}).get(output_type)


def run_output_format_matcher(client_code):
    """Main: download templates, analyse, create specs, upload.

    Returns the output-formats.json dict.
    """
    config = load_client_config(client_code)
    if not config:
        log.error(f"Client config not found for: {client_code}")
        return None

    templates_folder = config.get('folders', {}).get('templates')
    if not templates_folder:
        log.error("No templates folder configured")
        return None

    learning_folder = config.get('folders', {}).get('learning')
    if not learning_folder:
        log.error("No learning folder configured")
        return None

    # Download files from templates folder
    files = drive_client.list_files(templates_folder)
    if not files:
        log.warning("No files found in templates folder")
        output = {
            'formats': {},
            'generated_at': datetime.now().isoformat(),
        }
        drive_client.upload_or_update_json(learning_folder, 'output-formats.json', output)
        return output

    log.info(f"Found {len(files)} template files")

    formats = {}
    with tempfile.TemporaryDirectory() as tmpdir:
        for f in files:
            log.info(f"Analysing: {f['name']}")
            local_path = os.path.join(tmpdir, f['name'])
            try:
                drive_client.download_file(f['id'], local_path)
                spec = analyse_template(local_path, f['name'], client_code)
                if spec:
                    # Use the format key from Sonnet's response, or derive from filename
                    format_key = spec.get('key', _derive_format_key(f['name']))
                    formats[format_key] = spec
                    log.info(f"  Spec created: {format_key}")
            except Exception as e:
                log.warning(f"  Failed to analyse {f['name']}: {e}")

    output = {
        'formats': formats,
        'generated_at': datetime.now().isoformat(),
    }

    # Upload to 07-LEARNING
    drive_client.upload_or_update_json(learning_folder, 'output-formats.json', output)
    log.info("Output formats uploaded to 07-LEARNING")

    # Print summary
    print(f"\n{'='*50}")
    print(f"Output Format Specs — {config.get('client_name', client_code)}")
    print(f"{'='*50}")
    print(f"Templates analysed: {len(files)}")
    print(f"Specs created:      {len(formats)}")
    for key, spec in formats.items():
        spec_type = spec.get('type', 'unknown')
        col_count = len(spec.get('columns', []))
        detail = f"{col_count} columns" if spec_type == 'spreadsheet' else spec_type
        print(f"  - {key}: {detail}")
    print(f"{'='*50}")

    return output


def _derive_format_key(filename):
    """Derive a format key from a filename.

    e.g. 'Invoice Tracker Template.xlsx' -> 'invoice_tracker'
    """
    stem = Path(filename).stem.lower()
    # Remove common suffixes
    for suffix in ['template', 'sample', 'example', 'blank', 'master']:
        stem = stem.replace(suffix, '')
    # Clean up
    key = stem.strip(' _-').replace(' ', '_').replace('-', '_')
    # Remove double underscores
    while '__' in key:
        key = key.replace('__', '_')
    return key.strip('_') or 'unnamed'


def main():
    parser = argparse.ArgumentParser(description='Skill 9: Output Format Matcher')
    parser.add_argument('client_code', help='Client code (e.g. GILM)')
    args = parser.parse_args()
    run_output_format_matcher(args.client_code)


if __name__ == '__main__':
    main()
