"""Skill 10: Enhanced File Classifier.

Wraps the base Haiku classifier with:
  - Hard confidence thresholds from routing-rules.json
  - Client correction pattern integration
  - OCR triggering for poor images (Skill 7)
  - Classification learning from corrections

Usage:
    python scripts/skills/enhanced_classifier.py <client_code> <file_id>
"""

import os
import sys
import json
import argparse
import tempfile
import logging
from pathlib import Path

logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')
log = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from scripts.utils import claude_client, drive_client, sheets_client
from scripts.utils.prompt_builder import load_base_prompt, load_client_config
from scripts.skills.ocr_preprocessor import (
    process_file as ocr_process, is_image_file, is_pdf_file, is_text_searchable_pdf,
)


def load_routing_rules():
    """Load classification config from routing-rules.json."""
    path = PROJECT_ROOT / 'config' / 'workflows' / 'routing-rules.json'
    with open(path, 'r', encoding='utf-8') as f:
        return json.load(f)


def load_classification_corrections(client_code):
    """Load wrong_classification corrections from client's Drive."""
    config = load_client_config(client_code)
    if not config:
        return []
    learning_folder = config.get('folders', {}).get('learning')
    if not learning_folder:
        return []
    data = drive_client.download_json(learning_folder, 'corrections.json')
    if not data:
        return []
    return [
        c for c in data.get('corrections', [])
        if c.get('category') == 'wrong_classification'
    ]


def build_classifier_prompt(base_prompt, corrections=None):
    """Build the full classifier prompt with correction context."""
    prompt = base_prompt
    if corrections:
        prompt += '\n\n# CLASSIFICATION CORRECTIONS FROM PAST ERRORS\n'
        prompt += 'Use these to avoid repeating past misclassifications:\n'
        for c in corrections:
            prompt += (
                f"- File with characteristics of '{c.get('original_value', '?')}' "
                f"was actually '{c.get('corrected_value', '?')}': "
                f"{c.get('cause', 'no details')}\n"
            )
    return prompt


def classify_file(
    client_code, file_id=None, local_path=None,
    filename=None, content_text=None
):
    """Classify a file with enhanced logic.

    Accepts either a Drive file_id (downloads it) or a local_path.

    Returns dict:
        {
            'classification': str,
            'confidence': float,
            'status': 'NEW' | 'SPOT_CHECK' | 'NEEDS_CLASSIFICATION',
            'summary': str,
            'suggested_workflow': str,
            'ocr_applied': bool,
        }
    """
    rules = load_routing_rules()
    thresholds = rules['confidence_thresholds']
    config = load_client_config(client_code)

    # Download file if needed
    tmpdir = None
    if file_id and not local_path:
        meta = drive_client.get_file_metadata(file_id)
        filename = filename or meta['name']
        tmpdir = tempfile.mkdtemp()
        local_path = os.path.join(tmpdir, filename)
        drive_client.download_file(file_id, local_path)
    elif local_path:
        filename = filename or os.path.basename(local_path)

    # Load base classifier prompt
    base_prompt = load_base_prompt('classifier.txt')

    # Load client corrections for classification
    corrections = load_classification_corrections(client_code)
    system_prompt = build_classifier_prompt(base_prompt, corrections)

    # Get file content for classification
    if content_text:
        file_content = content_text
    else:
        file_content = _extract_text_content(local_path)

    # First classification attempt
    result = _run_classification(
        system_prompt, filename, file_content, local_path, client_code,
    )
    ocr_applied = False

    # Check confidence against thresholds
    confidence = result.get('confidence', 0)

    if confidence < thresholds['spot_check']:
        # Low confidence — check if OCR can help
        if local_path and (is_image_file(local_path) or
                           (is_pdf_file(local_path) and not is_text_searchable_pdf(local_path))):
            log.info(f"Low confidence ({confidence:.2f}), triggering OCR pre-processor")
            ocr_result = ocr_process(local_path)

            if ocr_result['ocr_applied'] and ocr_result['text']:
                # Re-classify with OCR text
                result = _run_classification(
                    system_prompt, filename, ocr_result['text'],
                    ocr_result['enhanced_path'], client_code,
                )
                confidence = result.get('confidence', 0)
                ocr_applied = True
                log.info(f"Post-OCR confidence: {confidence:.2f}")

    # Determine status
    if confidence >= thresholds['auto_route']:
        status = 'NEW'
    elif confidence >= thresholds['spot_check']:
        status = 'SPOT_CHECK'
    else:
        status = 'NEEDS_CLASSIFICATION'

    result['status'] = status
    result['ocr_applied'] = ocr_applied

    # Map classification to workflow
    classification = result.get('classification', 'OTHER')
    class_map = rules.get('classification_map', {})
    if classification in class_map and class_map[classification].get('handler'):
        result['suggested_workflow'] = class_map[classification]['handler']
    elif 'suggested_workflow' not in result:
        result['suggested_workflow'] = None

    # Log to tracking sheet
    if config and config.get('tracking_sheet_id'):
        try:
            from datetime import datetime
            sheets_client.append_row(
                config['tracking_sheet_id'],
                [
                    filename,
                    datetime.now().strftime('%Y-%m-%d %H:%M'),
                    classification,
                    f"{confidence:.2f}",
                    status,
                    result.get('suggested_workflow', ''),
                    '',  # QA Score
                    '',  # QA Status
                    f"OCR: {ocr_applied}" if ocr_applied else '',
                    '',  # Completed Date
                ],
            )
        except Exception as e:
            log.warning(f"Could not update tracking sheet: {e}")

    # Cleanup temp dir
    if tmpdir:
        import shutil
        shutil.rmtree(tmpdir, ignore_errors=True)

    return result


def _run_classification(system_prompt, filename, content, filepath, client_code):
    """Run a single classification call to Claude."""
    images = None
    # For images, send the image directly to Claude for vision classification
    if filepath and is_image_file(filepath):
        images = [claude_client.encode_image_file(filepath)]

    prompt = f"Filename: {filename}\n\nContent:\n{content[:3000]}"

    try:
        return claude_client.send_message_json(
            prompt=prompt,
            system_prompt=system_prompt,
            model_tier='MODEL_CLASSIFIER',
            max_tokens=300,
            client_code=client_code,
            workflow='classification',
            images=images,
        )
    except (json.JSONDecodeError, ValueError) as e:
        log.error(f"Classification response parse error: {e}")
        return {
            'classification': 'OTHER',
            'confidence': 0.0,
            'summary': 'Failed to parse classification response',
            'suggested_workflow': None,
        }


def _extract_text_content(filepath):
    """Extract text content from a file for classification."""
    if not filepath or not os.path.exists(filepath):
        return ''

    ext = Path(filepath).suffix.lower()

    if ext == '.pdf':
        try:
            import pdfplumber
            with pdfplumber.open(filepath) as pdf:
                text = ''
                for page in pdf.pages[:3]:
                    text += (page.extract_text() or '') + '\n'
                return text[:3000]
        except Exception:
            return ''

    elif ext in ('.doc', '.docx'):
        try:
            from docx import Document
            doc = Document(filepath)
            text = '\n'.join(p.text for p in doc.paragraphs)
            return text[:3000]
        except Exception:
            return ''

    elif ext in ('.txt', '.csv'):
        try:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read(3000)
        except Exception:
            return ''

    elif ext in ('.xlsx', '.xls'):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            ws = wb.active
            rows = []
            for row in ws.iter_rows(max_row=20, values_only=True):
                rows.append(' | '.join(str(c) for c in row if c))
            return '\n'.join(rows)[:3000]
        except Exception:
            return ''

    return ''


def main():
    parser = argparse.ArgumentParser(description='Skill 10: Enhanced File Classifier')
    parser.add_argument('client_code', help='Client code (e.g. GILM)')
    parser.add_argument('file_id', nargs='?', help='Google Drive file ID')
    parser.add_argument('--local', help='Local file path instead of Drive')
    args = parser.parse_args()

    result = classify_file(
        client_code=args.client_code,
        file_id=args.file_id,
        local_path=args.local,
    )

    print(f"\n{'='*50}")
    print(f"Enhanced Classification Result")
    print(f"{'='*50}")
    print(f"Classification: {result.get('classification')}")
    print(f"Confidence:     {result.get('confidence', 0):.2f}")
    print(f"Status:         {result.get('status')}")
    print(f"Workflow:       {result.get('suggested_workflow')}")
    print(f"OCR Applied:    {result.get('ocr_applied')}")
    print(f"Summary:        {result.get('summary', '')}")
    print(f"{'='*50}")


if __name__ == '__main__':
    main()
