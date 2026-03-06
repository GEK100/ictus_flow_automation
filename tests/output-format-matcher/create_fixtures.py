"""Generate test fixture files for Output Format Matcher tests.

Run once to create the test Excel and Word templates:
    python tests/output-format-matcher/create_fixtures.py
"""

import os
from pathlib import Path

FIXTURE_DIR = Path(__file__).parent


def create_invoice_tracker_excel():
    """Test A: Invoice tracker Excel with 14 columns + conditional formatting."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill, Font
    from datetime import date

    wb = Workbook()
    ws = wb.active
    ws.title = 'Invoices'
    ws.freeze_panes = 'A2'

    headers = [
        'Date', 'Invoice No', 'Supplier', 'Description', 'Net Amount',
        'VAT Rate', 'VAT Amount', 'Gross Amount', 'Category', 'Project',
        'Payment Due', 'Paid Date', 'Status', 'Notes',
    ]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)

    # Sample data rows
    rows = [
        [date(2025, 1, 15), 'INV-001', 'Acme Supplies', 'Bricks', 1000, 0.20, 200, 1200,
         'Materials', 'PRJ-101', date(2025, 2, 15), None, 'Unpaid', ''],
        [date(2025, 1, 20), 'INV-002', 'Quick Fix Ltd', 'Plumbing', 500, 0.20, 100, 600,
         'Subcontractor', 'PRJ-101', date(2025, 2, 20), date(2025, 2, 10), 'Paid', ''],
        [date(2025, 2, 1), 'INV-003', 'Steel Co', 'Steel beams', 3500, 0.20, 700, 4200,
         'Materials', 'PRJ-102', date(2025, 3, 1), None, 'Overdue', 'Chase required'],
    ]
    for r_idx, row in enumerate(rows, 2):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    # Number formats
    for row in range(2, 5):
        ws.cell(row=row, column=1).number_format = 'DD/MM/YYYY'
        ws.cell(row=row, column=5).number_format = '#,##0.00'
        ws.cell(row=row, column=6).number_format = '0%'
        ws.cell(row=row, column=7).number_format = '#,##0.00'
        ws.cell(row=row, column=8).number_format = '#,##0.00'
        ws.cell(row=row, column=11).number_format = 'DD/MM/YYYY'
        ws.cell(row=row, column=12).number_format = 'DD/MM/YYYY'

    # VAT formula: =E{row}*F{row}
    for row in range(2, 5):
        ws.cell(row=row, column=7, value=f'=E{row}*F{row}')
    # Gross formula: =E{row}+G{row}
    for row in range(2, 5):
        ws.cell(row=row, column=8, value=f'=E{row}+G{row}')

    # Conditional formatting: highlight overdue in red
    red_fill = PatternFill(start_color='FFCCCC', end_color='FFCCCC', fill_type='solid')
    ws.conditional_formatting.add('M2:M1000',
        CellIsRule(operator='equal', formula=['"Overdue"'], fill=red_fill))

    # Green for paid
    green_fill = PatternFill(start_color='CCFFCC', end_color='CCFFCC', fill_type='solid')
    ws.conditional_formatting.add('M2:M1000',
        CellIsRule(operator='equal', formula=['"Paid"'], fill=green_fill))

    # Column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['H'].width = 14
    ws.column_dimensions['M'].width = 12

    path = FIXTURE_DIR / 'test_a_invoice_tracker.xlsx'
    wb.save(path)
    print(f"Created: {path}")


def create_letterhead_docx():
    """Test B: Word letterhead template."""
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # Header
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = 'Gilmartins Property Maintenance'
    hp.style.font.size = Pt(14)
    hp.style.font.bold = True

    # Footer
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.text = 'Tel: 01onal 123456 | info@gilmartins.co.uk | www.gilmartins.co.uk'
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Body content
    doc.add_heading('Subject Line Here', level=1)
    doc.add_paragraph('Dear [Name],')
    doc.add_paragraph(
        'This is a template letter for correspondence. It demonstrates the '
        'standard layout including margins, fonts, and structure.'
    )
    doc.add_paragraph('Yours sincerely,')
    doc.add_paragraph('[Sender Name]')

    path = FIXTURE_DIR / 'test_b_letterhead.docx'
    doc.save(path)
    print(f"Created: {path}")


def create_csv_template():
    """Test C: Simple CSV for validation testing."""
    import csv

    path = FIXTURE_DIR / 'test_c_timesheet.csv'
    with open(path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(['Date', 'Employee', 'Project', 'Hours', 'Rate', 'Total', 'Approved'])
        writer.writerow(['15/01/2025', 'John Smith', 'PRJ-101', '8', '25.00', '200.00', 'Yes'])
        writer.writerow(['16/01/2025', 'John Smith', 'PRJ-101', '7.5', '25.00', '187.50', 'Yes'])
        writer.writerow(['17/01/2025', 'Jane Doe', 'PRJ-102', '8', '30.00', '240.00', 'No'])

    print(f"Created: {path}")


def create_expected_outputs():
    """Create expected output JSON files for test validation."""
    import json

    # Test A expected: key fields that must be present
    test_a_expected = {
        'type': 'spreadsheet',
        'expected_columns': [
            'Date', 'Invoice No', 'Supplier', 'Description', 'Net Amount',
            'VAT Rate', 'VAT Amount', 'Gross Amount', 'Category', 'Project',
            'Payment Due', 'Paid Date', 'Status', 'Notes',
        ],
        'expected_column_count': 14,
        'expected_types': {
            'Date': 'date',
            'Net Amount': 'currency',
            'VAT Rate': 'number',
            'Status': 'text',
        },
        'expected_conditional_formatting': True,
        'expected_frozen_panes': True,
    }

    # Test B expected
    test_b_expected = {
        'type': 'document',
        'expected_has_header': True,
        'expected_has_footer': True,
        'expected_orientation': 'portrait',
        'expected_margins_present': True,
    }

    # Test C expected
    test_c_expected = {
        'type': 'spreadsheet',
        'expected_columns': ['Date', 'Employee', 'Project', 'Hours', 'Rate', 'Total', 'Approved'],
        'expected_column_count': 7,
        'expected_types': {
            'Date': 'date',
            'Hours': 'number',
            'Rate': 'number',
            'Total': 'number',
        },
    }

    for name, data in [
        ('test_a_expected.json', test_a_expected),
        ('test_b_expected.json', test_b_expected),
        ('test_c_expected.json', test_c_expected),
    ]:
        path = FIXTURE_DIR / name
        with open(path, 'w') as f:
            json.dump(data, f, indent=2)
        print(f"Created: {path}")


if __name__ == '__main__':
    create_invoice_tracker_excel()
    create_letterhead_docx()
    create_csv_template()
    create_expected_outputs()
    print("\nAll test fixtures created.")
